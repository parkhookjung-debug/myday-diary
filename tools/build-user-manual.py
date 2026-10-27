"""Build the Korean user manual from a shared documentation source. Requires python-docx."""
from pathlib import Path
import argparse,json,re,os
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from PIL import Image

def inline(paragraph,text):
    for i,value in enumerate(re.split(r'(\*\*.*?\*\*)',text)):
        if not value:continue
        r=paragraph.add_run(value[2:-2] if value.startswith('**') else value);r.bold=value.startswith('**')

def external_link(paragraph,label,url,anchor=None):
    h=OxmlElement('w:hyperlink')
    if anchor:h.set(qn('w:anchor'),anchor)
    else:
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
        h.set(qn('r:id'),paragraph.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r');rp=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'000000');rp.append(color)
    underline=OxmlElement('w:u');underline.set(qn('w:val'),'single');rp.append(underline);r.append(rp)
    text=OxmlElement('w:t');text.text=label;r.append(text);h.append(r);paragraph._p.append(h)

def field(paragraph,instruction,result=''):
    run=paragraph.add_run();a=OxmlElement('w:fldChar');a.set(qn('w:fldCharType'),'begin');run._r.append(a)
    text=OxmlElement('w:instrText');text.set(qn('xml:space'),'preserve');text.text=instruction;run._r.append(text)
    sep=OxmlElement('w:fldChar');sep.set(qn('w:fldCharType'),'separate');run._r.append(sep)
    paragraph.add_run(result);end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end');paragraph.add_run()._r.append(end)

def render(content,root,docx_path,md_path):
    root=Path(root);d=Document();sec=d.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=Inches(.72);sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.76)
    sec.header_distance=sec.footer_distance=Inches(.28)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
        st=d.styles[name];st.font.name='Malgun Gothic';st.font.color.rgb=RGBColor(0,0,0)
        fonts=st.element.get_or_add_rPr().get_or_add_rFonts()
        for theme in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
            fonts.attrib.pop(qn('w:'+theme),None)
        for script in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+script),'Malgun Gothic')
    normal=d.styles['Normal'];normal.font.size=Pt(11);normal.paragraph_format.line_spacing=1.18;normal.paragraph_format.space_after=Pt(6)
    normal.paragraph_format.widow_control=True
    for name,size in [('Title',28),('Subtitle',12),('Heading 1',19),('Heading 2',13),('Heading 3',11.5)]:
        st=d.styles[name];st.font.size=Pt(size);st.font.bold=name!='Subtitle';st.paragraph_format.space_before=Pt(15 if name!='Title' else 0);st.paragraph_format.space_after=Pt(10);st.paragraph_format.keep_with_next=True
    d.styles['Heading 1'].paragraph_format.page_break_before=False
    d.styles['Caption'].font.size=Pt(9);d.styles['Caption'].paragraph_format.space_after=Pt(12)
    for name in ['Title','Subtitle','Heading 1','Heading 2','Heading 3']:
        pr=d.styles[name].element.find(qn('w:pPr'))
        if pr is not None:
            for border in list(pr.findall(qn('w:pBdr'))):pr.remove(border)
    header=sec.header.paragraphs[0];header.text='MyDay 0.12.0 사용자 설명서';header.runs[0].font.size=Pt(8)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    field(footer,'PAGE','1');footer.add_run(' / ');field(footer,'NUMPAGES','1')
    for r in footer.runs:r.font.size=Pt(8)
    d.core_properties.title=content['title'];d.core_properties.subject='MyDay 0.12.0 설치와 기능과 사용 절차';d.core_properties.author='MyDay project documentation'
    d.add_paragraph(content['title'],'Title');d.add_paragraph('Windows Mac Android 사용자와 공동 작업자를 위한 안내','Subtitle')
    for line in ['앱 버전 0.12.0','문서 기준일 2026년 10월 28일','기준 소스 cc358ae 작업 브랜치 codex/windows-app']:
        d.add_paragraph(line)
    d.add_paragraph('이 설명서는 설치부터 일기 한 편 작성, 사진과 화면 꾸미기, 상몬 성장과 소품 수집, 기록 보관과 복구까지 직접 따라 할 수 있도록 정리한 사용자 매뉴얼입니다.')
    d.add_paragraph('순서대로 처음 사용할 때는 본문을 따라 읽고, 기능을 찾을 때는 목차와 부록의 바로 찾기 표를 사용하세요. 화면 예시의 개인 기록은 테스트 자료이며 실제 개인 일기나 PIN을 포함하지 않습니다.')
    d.add_paragraph('현재 지원하는 기능과 지원하지 않는 기능을 구분했습니다. 운영체제별 조작·저장 방식과 데이터가 바뀌는 범위는 관련 절차에 함께 설명했습니다.')
    d.add_page_break();toc_title=d.add_paragraph('목차','Heading 1');toc_title.paragraph_format.page_break_before=False
    chapters=[n for n in content['nodes'] if n['type']=='chapter']
    for i,node in enumerate(chapters,1):
        entry=d.add_paragraph();entry.paragraph_format.space_after=Pt(2);entry.paragraph_format.line_spacing=1.05
        entry.paragraph_format.tab_stops.add_tab_stop(Inches(6.45),2,1)
        external_link(entry,f'{i:02d}  '+node['text'],'',anchor=f'chapter_{i}')
        entry.add_run('\t');field(entry,'PAGEREF chapter_'+str(i)+' \\h','')
        for run in entry.runs:run.font.size=Pt(10.5)
    md=['# '+content['title'],'','앱 0.12.0 · 문서 기준일 2026-10-28 · 소스 cc358ae','',
        '설치부터 기록·꾸미기·성장·백업·문제 해결까지 설명합니다. Windows·Mac·Android의 차이와 현재 지원 범위를 함께 확인하세요.','',
        '[편집 가능한 Word 설명서](manual/MyDay-0.12.0-User-Manual.docx)','', '## 목차','']
    for i,node in enumerate(chapters,1):md.append(f'- [{i:02d} {node["text"]}](#chapter-{i})')
    md.append('');chapter=0;bookmark=1
    for n in content['nodes']:
        kind=n['type']
        if kind=='chapter':
            chapter+=1;title=f'{chapter:02d} '+n['text'];par=d.add_paragraph(title,'Heading 1')
            if chapter in (1,36):par.paragraph_format.page_break_before=True
            a=OxmlElement('w:bookmarkStart');a.set(qn('w:id'),str(bookmark));a.set(qn('w:name'),f'chapter_{chapter}');par._p.insert(0,a)
            b=OxmlElement('w:bookmarkEnd');b.set(qn('w:id'),str(bookmark));par._p.append(b);bookmark+=1
            md.extend([f'<a id="chapter-{chapter}"></a>','', '## '+title,''])
        elif kind=='sub':d.add_paragraph(n['text'],'Heading 2');md.extend(['### '+n['text'],''])
        elif kind=='p':inline(d.add_paragraph(),n['text']);md.extend([n['text'],''])
        elif kind in ('steps','bullets'):
            for i,line in enumerate(n['items'],1):
                par=d.add_paragraph();par.paragraph_format.left_indent=Inches(.18);par.paragraph_format.first_line_indent=Inches(-.18)
                prefix=f'{i}. ' if kind=='steps' else '• ';par.add_run(prefix);inline(par,line)
                md.append((f'{i}. ' if kind=='steps' else '- ')+line)
            md.append('')
        elif kind=='table':
            heads=n['headers'];rows=n['rows'];table=d.add_table(rows=1,cols=len(heads));table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
            borders=OxmlElement('w:tblBorders')
            for edge in ['top','left','bottom','right','insideH','insideV']:
                border=OxmlElement('w:'+edge);border.set(qn('w:val'),'single');border.set(qn('w:sz'),'4');border.set(qn('w:color'),'D9D9D9');borders.append(border)
            table._tbl.tblPr.append(borders)
            ratios={2:[.33,.67],3:[.24,.50,.26],4:[.21,.26,.265,.265],5:[.25,.15,.13,.235,.235]}[len(heads)]
            for col,ratio in zip(table.columns,ratios):col.width=Inches(6.98*ratio)
            for i,text in enumerate(heads):table.rows[0].cells[i].text=text
            repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for row in rows:
                cells=table.add_row().cells
                for i,text in enumerate(row):cells[i].text=str(text)
            for ri,row in enumerate(table.rows):
                pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
                for ci,cell in enumerate(row.cells):
                    cell.width=Inches(6.98*ratios[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    cp=cell._tc.get_or_add_tcPr();marg=OxmlElement('w:tcMar')
                    for edge,value in [('top',72),('bottom',72),('left',95),('right',95)]:
                        e=OxmlElement('w:'+edge);e.set(qn('w:w'),str(value));e.set(qn('w:type'),'dxa');marg.append(e)
                    cp.append(marg)
                    if ri==0:
                        sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E9ECEF');cp.append(sh)
                    for par in cell.paragraphs:
                        par.paragraph_format.space_after=Pt(1);par.paragraph_format.space_before=Pt(1);par.paragraph_format.line_spacing=1.14
                        for run in par.runs:run.font.size=Pt(10.5);run.bold=ri==0
            d.add_paragraph().paragraph_format.space_after=Pt(2)
            esc=lambda value:str(value).replace('|','/').replace('\n','<br>')
            md.append('| '+' | '.join(map(esc,heads))+' |');md.append('| '+' | '.join('---' for _ in heads)+' |')
            md.extend('| '+' | '.join(map(esc,row))+' |' for row in rows);md.append('')
        elif kind=='figure':
            src=root/n['path'];assert src.exists(),src
            w,h=Image.open(src).size;max_h=n.get('height',4);width=min(6.65,max_h*w/h)
            par=d.add_paragraph();par.alignment=WD_ALIGN_PARAGRAPH.CENTER;par.paragraph_format.keep_with_next=True
            par.add_run().add_picture(str(src),width=Inches(width))
            par=d.add_paragraph(n['caption'],'Caption');par.alignment=WD_ALIGN_PARAGRAPH.CENTER
            relative=os.path.relpath(src,md_path.parent).replace('\\','/');md.extend([f'![{n["caption"]}]({relative})',''])
        elif kind=='code':
            for line in n['text'].splitlines():
                par=d.add_paragraph();par.paragraph_format.space_after=Pt(2)
                run=par.add_run(line);run.font.name='Consolas';run.font.size=Pt(9)
            md.extend(['```text',n['text'],'```',''])
        elif kind=='link':
            par=d.add_paragraph();external_link(par,n['label'],n['url']);md.extend([f'[{n["label"]}]({n["url"]})',''])
        else:raise ValueError(kind)
    settings=d.settings.element;update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
    docx_path.parent.mkdir(parents=True,exist_ok=True);d.save(docx_path);md_path.write_text('\n'.join(md),encoding='utf-8',newline='\n')
    print(json.dumps({'chapters':len(chapters),'nodes':len(content['nodes']),'docx':str(docx_path),'markdown':str(md_path)},ensure_ascii=False))

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--root',type=Path,default=Path(__file__).resolve().parents[1]);args=parser.parse_args()
    root=args.root;content=json.loads((root/'docs/manual/user-manual-content.json').read_text(encoding='utf-8'))
    render(content,root,root/'docs/manual/MyDay-0.12.0-User-Manual.docx',root/'docs/USER-MANUAL.md')
